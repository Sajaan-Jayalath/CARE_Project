"""Write applied fixes into a copy of the uploaded document. The original is never changed.

Each edit addresses an original extracted section by character range. Writers map
that range back onto the file's own structure (Word runs, slide runs, spreadsheet
cells, PDF text blocks, plain text offsets). An edit whose source text no longer
matches, or whose location a format cannot edit, is reported as not applied rather
than guessed.
"""
import csv
from collections import defaultdict
from io import BytesIO, StringIO
from zipfile import ZipFile

from lxml import etree

from backend.database.database import BackendError
from backend.extraction.office import A, S, W

XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
PARSER = etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False)
MEDIA = {
    "pdf": "application/pdf", "txt": "text/plain", "text": "text/plain", "md": "text/markdown",
    "csv": "text/csv", "tsv": "text/tab-separated-values",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "xlsm": "application/vnd.ms-excel.sheet.macroEnabled.12",
}


def corrected_document(store, job_id):
    """Return (bytes, filename, media type, number of edits not applied)."""
    job = store.job(job_id)
    document = store.document(job["upload_id"], original=True)
    data, unapplied = corrected_bytes(document, store.edits(job_id))
    stem, dot, ext = document["filename"].rpartition(".")
    name = f"{stem or document['filename']} (CARE revised){dot}{ext if stem else ''}"
    return data, name, MEDIA[document["extraction"]["format"]], unapplied


def corrected_bytes(document, edits):
    """The upload's bytes with the edits written in, and how many edits could not be."""
    extraction = document["extraction"]
    writer = WRITERS.get(extraction["format"])
    if writer is None:
        raise BackendError("export_unsupported", "A corrected copy cannot be created for this file type. Save it as .xlsx and upload it again.", 422)
    sections = {s["section_id"]: s["text"] for s in (extraction.get("document") or {}).get("sections", [])}
    by_section, skipped = defaultdict(list), 0
    for edit in edits:
        if sections.get(edit["section_id"], "")[edit["start"]:edit["end"]] == edit["original"]:
            by_section[edit["section_id"]].append(edit)
        else:
            skipped += 1
    data, unapplied = writer(document["original"], extraction, sections, dict(by_section))
    return data, skipped + unapplied


def _edits(edits):
    return [(e["start"], e["end"], e["replacement"]) for e in edits]


def _apply(text, edits):
    for start, end, replacement in sorted(edits, reverse=True):
        text = text[:start] + replacement + text[end:]
    return text


# Plain text, Markdown, pasted text -----------------------------------------
def write_text(original, extraction, sections, by_section):
    utf16 = original.startswith((b"\xff\xfe", b"\xfe\xff"))
    bom = original.startswith(b"\xef\xbb\xbf")
    text = original.decode("utf-16" if utf16 else "utf-8-sig")
    edits = []
    for sid, items in by_section.items():
        base = extraction["sources"][sid].get("input_start_char", 0)
        edits += [(base + s, base + e, r) for s, e, r in _edits(items)]
    text = _apply(text, edits)
    return text.encode("utf-16" if utf16 else "utf-8-sig" if bom else "utf-8"), 0


# CSV / TSV ------------------------------------------------------------------
def write_table(delimiter):
    def writer(original, extraction, sections, by_section):
        utf16 = original.startswith((b"\xff\xfe", b"\xfe\xff"))
        text = original.decode("utf-16" if utf16 else "utf-8-sig")
        rows = [list(r) for r in csv.reader(StringIO(text, newline=""), delimiter=delimiter, strict=True)]
        for sid, items in by_section.items():
            source = extraction["sources"][sid]
            row, col = source["row"] - 1, source["column"] - 1
            rows[row][col] = _apply(rows[row][col], _edits(items))
        out = StringIO(newline="")
        csv.writer(out, delimiter=delimiter, lineterminator="\r\n" if "\r\n" in text else "\n").writerows(rows)
        return out.getvalue().encode("utf-16" if utf16 else "utf-8"), 0
    return writer


# Shared run editing ---------------------------------------------------------
def _replace_in_pieces(pieces, start, end, replacement):
    """pieces: (element holding text or None for a 1-char break, text, break element).

    The replacement lands in the first text run the range touches (keeping its
    formatting); the rest of the range is trimmed from later runs and breaks.
    """
    pos, anchor = 0, None
    for holder, text, brk in pieces:
        ps, pe = pos, pos + len(text)
        pos = pe
        if not (ps < end and pe > start):
            continue
        a, b = max(start, ps) - ps, min(end, pe) - ps
        if holder is None:
            brk.getparent().remove(brk)
            continue
        if anchor is None:
            holder.text, anchor = text[:a] + replacement + text[b:], holder
        else:
            holder.text = text[:a] + text[b:]
        holder.set(XML_SPACE, "preserve")
    return anchor is not None


def _apply_runs(get_pieces, items):
    """Apply edits right-to-left, rebuilding the run map after each change."""
    applied = 0
    for start, end, replacement in sorted(_edits(items), reverse=True):
        applied += _replace_in_pieces(get_pieces(), start, end, replacement)
    return len(items) - applied


def _children(node):
    return [c for c in node if isinstance(c.tag, str)]


def _zip_copy(original, replaced):
    out = BytesIO()
    with ZipFile(BytesIO(original)) as zin, ZipFile(out, "w") as zout:
        for item in zin.infolist():
            zout.writestr(item, replaced.get(item.filename) or zin.read(item.filename))
    return out.getvalue()


def _serialize(root):
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


# Word -----------------------------------------------------------------------
def _word_pieces(paragraph):
    """Mirror docx_extractor.paragraph_text: w:t text, tabs, breaks; skip nested paragraphs."""
    pieces = []
    def walk(element):
        if element is not paragraph and element.tag == W + "p":
            return
        if element.tag in (W + "del", W + "instrText"):
            return
        if element.tag == W + "t":
            pieces.append((element, element.text or "", None))
        elif element.tag == W + "tab":
            pieces.append((None, "\t", element))
        elif element.tag in (W + "br", W + "cr"):
            pieces.append((None, "\n", element))
        else:
            for child in _children(element):
                walk(child)
    walk(paragraph)
    return pieces


def write_docx(original, extraction, sections, by_section):
    parts, unapplied = {}, 0
    with ZipFile(BytesIO(original)) as archive:
        for sid, items in by_section.items():
            source = extraction["sources"][sid]
            part = source.get("part")
            if part not in parts:
                parts[part] = etree.fromstring(archive.read(part), PARSER)
            root = parts[part]
            if source.get("kind") == "paragraph":
                node = root
                for index in source["xml_path"]:
                    node = _children(node)[index - 1]
                get_pieces = lambda node=node: _word_pieces(node)
            elif source.get("kind") == "drawing_text":
                node = list(root.iter(A + "p"))[source["drawing_paragraph"] - 1]
                get_pieces = lambda node=node: [(t, t.text or "", None) for t in node.iter(A + "t")]
            else:
                unapplied += len(items)
                continue
            if "".join(p[1] for p in get_pieces()) != sections[sid]:
                unapplied += len(items)
                continue
            unapplied += _apply_runs(get_pieces, items)
        return _zip_copy(original, {part: _serialize(root) for part, root in parts.items()}), unapplied


# PowerPoint -----------------------------------------------------------------
def _slide_pieces(paragraph):
    """Mirror python-pptx paragraph.text: run text, line breaks as \\v, field text."""
    pieces = []
    for child in _children(paragraph):
        if child.tag in (A + "r", A + "fld"):
            t = child.find(A + "t")
            if t is not None:
                pieces.append((t, t.text or "", None))
        elif child.tag == A + "br":
            pieces.append((None, "\v", child))
    return pieces


def _find_shape(shapes, shape_id):
    for shape in shapes:
        if shape.shape_id == shape_id:
            return shape
        if hasattr(shape, "shapes"):
            found = _find_shape(shape.shapes, shape_id)
            if found is not None:
                return found
    return None


def write_pptx(original, extraction, sections, by_section):
    from pptx import Presentation
    deck = Presentation(BytesIO(original))
    slides = list(deck.slides)
    unapplied = 0
    for sid, items in by_section.items():
        source = extraction["sources"][sid]
        slide = slides[source["slide"] - 1]
        kind = source.get("kind")
        paragraphs = None
        if kind == "speaker_notes":
            paragraphs = [slide.notes_slide.notes_text_frame.paragraphs[source["paragraph"] - 1]]
        elif kind in ("slide_paragraph", "table_cell"):
            shape = _find_shape(slide.shapes, source["shape_id"])
            if shape is not None and kind == "slide_paragraph":
                paragraphs = [shape.text_frame.paragraphs[source["paragraph"] - 1]]
            elif shape is not None:
                paragraphs = list(shape.table.cell(source["row"] - 1, source["column"] - 1).text_frame.paragraphs)
        if not paragraphs or "\n".join(p.text for p in paragraphs) != sections[sid]:
            unapplied += len(items)
            continue
        # Table cells join paragraphs with \n; a fix never crosses a line break.
        offsets, offset = [], 0
        for paragraph in paragraphs:
            offsets.append((offset, paragraph))
            offset += len(paragraph.text) + 1
        groups = defaultdict(list)
        for edit in items:
            index = next((i for i, (o, p) in enumerate(offsets) if o <= edit["start"] and edit["end"] <= o + len(p.text)), None)
            if index is None:
                unapplied += 1
                continue
            o = offsets[index][0]
            groups[index].append(dict(edit, start=edit["start"] - o, end=edit["end"] - o))
        for index, local in groups.items():
            unapplied += _apply_runs(lambda p=offsets[index][1]._p: _slide_pieces(p), local)
    out = BytesIO()
    deck.save(out)
    return out.getvalue(), unapplied


# Excel ----------------------------------------------------------------------
def write_xlsx(original, extraction, sections, by_section):
    parts, unapplied = {}, 0
    with ZipFile(BytesIO(original)) as archive:
        names = set(archive.namelist())
        shared = etree.fromstring(archive.read("xl/sharedStrings.xml"), PARSER) if "xl/sharedStrings.xml" in names else None
        for sid, items in by_section.items():
            source = extraction["sources"][sid]
            part = source.get("part")
            if source.get("kind") != "cell" or source.get("cached_formula") or part not in names:
                unapplied += len(items)
                continue
            if part not in parts:
                parts[part] = etree.fromstring(archive.read(part), PARSER)
            cell = next((c for c in parts[part].iter(S + "c") if c.get("r") == source["cell"]), None)
            kind = cell.get("t", "n") if cell is not None else None
            if kind == "s" and shared is not None:
                strings = [c for c in shared if c.tag == S + "si"]
                value = cell.find(S + "v")
                current = "".join(t.text or "" for t in strings[int(value.text)].iter(S + "t"))
                if current != sections[sid]:
                    unapplied += len(items)
                    continue
                # Shared strings can be reused by other cells: give this cell its own copy.
                item = etree.SubElement(shared, S + "si")
                t = etree.SubElement(item, S + "t")
                t.text = _apply(current, _edits(items))
                t.set(XML_SPACE, "preserve")
                value.text = str(len(strings))
                shared.set("uniqueCount", str(len(strings) + 1))
            elif kind == "inlineStr":
                current = "".join(t.text or "" for t in cell.iter(S + "t"))
                if current != sections[sid]:
                    unapplied += len(items)
                    continue
                for old in cell.findall(S + "is"):
                    cell.remove(old)
                inline = etree.SubElement(cell, S + "is")
                t = etree.SubElement(inline, S + "t")
                t.text = _apply(current, _edits(items))
                t.set(XML_SPACE, "preserve")
            else:
                unapplied += len(items)
        replaced = {part: _serialize(root) for part, root in parts.items()}
        if shared is not None:
            replaced["xl/sharedStrings.xml"] = _serialize(shared)
        return _zip_copy(original, replaced), unapplied


# PDF ------------------------------------------------------------------------
def _base_font(flags, font):
    bold = bool(flags & 16) or "bold" in font.lower()
    if flags & 8:
        return "cobo" if bold else "cour"
    if flags & 4:
        return "tibo" if bold else "tiro"
    return "hebo" if bold else "helv"


def write_pdf(original, extraction, sections, by_section):
    """Replace each edited text block in place: remove its text, write the corrected text in its box."""
    import pymupdf
    unapplied = 0
    with pymupdf.open(stream=original, filetype="pdf") as pdf:
        plans = defaultdict(list)
        for sid, items in by_section.items():
            source = extraction["sources"][sid]
            if not source.get("bbox") or not source.get("page"):
                unapplied += len(items)
                continue
            page = pdf[source["page"] - 1]
            rect = pymupdf.Rect(source["bbox"])
            sizes = defaultdict(int)
            style = None
            for block in page.get_text("dict", clip=rect)["blocks"]:
                for line in block.get("lines", []):
                    for span in line["spans"]:
                        sizes[round(span["size"], 1)] += len(span["text"])
                        style = style or span
            if style is None:
                unapplied += len(items)
                continue
            text = " ".join(_apply(sections[sid], _edits(items)).split())
            plans[source["page"] - 1].append((rect, text, max(sizes, key=sizes.get), style))
        for index, blocks in plans.items():
            page = pdf[index]
            for rect, *_ in blocks:
                page.add_redact_annot(rect, fill=False)
            page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE, graphics=pymupdf.PDF_REDACT_LINE_ART_NONE)
            for rect, text, size, style in blocks:
                color = tuple(((style["color"] >> s) & 255) / 255 for s in (16, 8, 0))
                font = _base_font(style["flags"], style["font"])
                box = pymupdf.Rect(rect.x0, rect.y0, rect.x1, rect.y1 + size * 0.6)
                fitted = False
                # Shrink the font until the corrected text fits the original block.
                for step in range(0, 60):
                    fontsize = size * (1 - step * 0.015)
                    shape = page.new_shape()
                    if shape.insert_textbox(box, text, fontsize=fontsize, fontname=font, color=color) >= 0:
                        shape.commit()
                        fitted = True
                        break
                if not fitted:
                    unapplied += 1
        return pdf.tobytes(garbage=3, deflate=True), unapplied


WRITERS = {"text": write_text, "txt": write_text, "md": write_text, "csv": write_table(","), "tsv": write_table("\t"),
           "docx": write_docx, "pptx": write_pptx, "xlsx": write_xlsx, "xlsm": write_xlsx, "pdf": write_pdf}
