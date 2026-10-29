"""Bounded in-memory image extraction, separate assets from their occurrences."""
import hashlib
import posixpath
from io import BytesIO
from pathlib import PurePosixPath
import warnings

from PIL import Image, UnidentifiedImageError
from backend.ai.schemas.image_schema import ImageOccurrence
from backend.extraction.common import ExtractionError
from backend.extraction.office import package, xml, W, A, R

MIME = {"PNG": "image/png", "JPEG": "image/jpeg", "TIFF": "image/tiff",
        "BMP": "image/bmp", "GIF": "image/gif", "WEBP": "image/webp"}


class ImageBatch:
    def __init__(self, out, settings):
        self.out, self.settings = out, settings
        self.assets = {}
        self.first = {}
        self.pixels = self.bytes = 0

    def warn(self, code, message, **location):
        self.out.warn(code, message, **location)

    def room(self):
        if len(self.out.images) >= self.settings.max_images:
            self.warn("image_count_limit", "Additional images were skipped at the configured image limit.")
            return False
        return True

    def add(self, data, kind="embedded", **location):
        if not self.room():
            return
        try:
            digest = hashlib.sha256(data).hexdigest()
            duplicate = self.first.get(digest)
            with warnings.catch_warnings():
                warnings.simplefilter("error", Image.DecompressionBombWarning)
                with Image.open(BytesIO(data)) as image:
                    width, height = image.size
                    mime = MIME.get(image.format)
                    if not mime:
                        raise ValueError("unsupported raster format")
                    if width * height > self.settings.max_pixels:
                        raise OverflowError()
                    if not duplicate:
                        if (self.pixels + width * height > self.settings.max_total_pixels or
                                self.bytes + len(data) > self.settings.max_total_bytes):
                            self.warn("image_total_limit", "Cumulative image memory limit reached.", **location)
                            return
                        image.load()
                    if getattr(image, "n_frames", 1) > 1:
                        self.warn("image_frames_skipped", "Only the first frame of this multi-frame image is assessed.", **location)
            identifier = f"{self.out.document_id}-image-{len(self.out.images) + 1}"
            context = self.context(location)
            occurrence = ImageOccurrence(document_id=self.out.document_id, image_id=identifier,
                asset_sha256=digest, image_index=len(self.out.images) + 1,
                source_type=self.out.format, kind=kind, mime_type=mime, width=width, height=height,
                page=location.get("page"), slide=location.get("slide"), location=location,
                surrounding_text="\n".join(s.text for s in context)[:8000],
                context_section_ids=[s.section_id for s in context], duplicate_of=duplicate)
            self.out.images.append(occurrence)
            if not duplicate:
                self.assets[digest] = data
                self.first[digest] = identifier
                self.pixels += width * height
                self.bytes += len(data)
            return occurrence
        except (OverflowError, Image.DecompressionBombWarning, Image.DecompressionBombError):
            self.warn("image_pixel_limit", "Image exceeds the safe decoded-pixel limit.", **location)
        except (UnidentifiedImageError, ValueError, OSError):
            self.warn("invalid_image", "Corrupted or unsupported image was skipped.", **location)

    def context(self, location):
        sections = self.out.sections
        if location.get("page"):
            sections = [s for s in sections if s.page == location["page"]]
        elif location.get("slide"):
            sections = [s for s in sections if s.slide == location["slide"]]
        elif location.get("part"):
            sections = [s for s in sections if self.out.sources[s.section_id].get("part") == location["part"]]
            if location.get("paragraph"):
                sections = [s for s in sections if s.paragraph and abs(s.paragraph - location["paragraph"]) <= 1]
        return sections[:30]


def pdf_images(data, batch):
    import pymupdf
    with pymupdf.open(stream=data, filetype="pdf") as pdf:
        for number, page in enumerate(pdf, 1):
            if not batch.room():
                break
            infos = page.get_image_info(xrefs=True)
            drawings = bool(page.get_drawings())
            if not infos and not drawings:
                continue
            # Whole-page rendering preserves scanned mosaics and vector diagrams.
            # Do not also OCR embedded assets from that same rendered page.
            if not page.get_text().strip() or drawings:
                width = int(page.rect.width * batch.settings.pdf_dpi / 72) + 1
                height = int(page.rect.height * batch.settings.pdf_dpi / 72) + 1
                if width * height > batch.settings.max_pixels or batch.pixels + width * height > batch.settings.max_total_pixels:
                    batch.warn("image_pixel_limit", "PDF page render exceeds the image pixel budget.", page=number)
                    continue
                try:
                    pix = page.get_pixmap(dpi=batch.settings.pdf_dpi, alpha=False)
                    batch.add(pix.tobytes("png"), kind="page_render", page=number,
                              bbox=list(page.rect), coordinate_system="PDF points, top-left origin",
                              render_dpi=batch.settings.pdf_dpi)
                except Exception:
                    batch.warn("image_render_failed", "PDF page could not be rendered.", page=number)
                continue
            for info in infos:
                if not batch.room():
                    break
                location = dict(page=number, bbox=list(info["bbox"]),
                                transform=list(info["transform"]), coordinate_system="PDF points, top-left origin")
                if info["width"] * info["height"] > batch.settings.max_pixels or batch.pixels + info["width"] * info["height"] > batch.settings.max_total_pixels:
                    batch.warn("image_pixel_limit", "PDF image exceeds the image pixel budget.", page=number)
                    continue
                try:
                    xref = info.get("xref")
                    if xref:
                        raw = pdf.extract_image(xref)
                        if raw.get("smask"):
                            pix = pymupdf.Pixmap(pdf, xref)
                            mask = pymupdf.Pixmap(pdf, raw["smask"])
                            blob = pymupdf.Pixmap(pix, mask).tobytes("png")
                        else:
                            blob = raw["image"]
                    else:
                        # Inline images cannot always be addressed as PDF assets.
                        rect = pymupdf.Rect(info["bbox"]) & page.rect
                        if rect.is_empty:
                            continue
                        rw = int(rect.width * batch.settings.pdf_dpi / 72) + 1
                        rh = int(rect.height * batch.settings.pdf_dpi / 72) + 1
                        if rw * rh > batch.settings.max_pixels or batch.pixels + rw * rh > batch.settings.max_total_pixels:
                            raise OverflowError()
                        blob = page.get_pixmap(clip=rect, dpi=batch.settings.pdf_dpi, alpha=False).tobytes("png")
                    batch.add(blob, **location)
                except Exception:
                    batch.warn("image_extraction_failed", "A PDF image could not be recovered.", **location)


def pptx_images(data, batch):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    deck = Presentation(BytesIO(data))
    for number, slide in enumerate(deck.slides, 1):
        def walk(items, parents=()):
            for shape in items:
                if not batch.room():
                    return
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    walk(shape.shapes, parents + (shape.shape_id,))
                elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    loc = dict(slide=number, shape_id=shape.shape_id, group_path=list(parents),
                               bbox=[shape.left, shape.top, shape.width, shape.height],
                               coordinate_system="shape-local EMU x,y,width,height",
                               crop=[shape.crop_left, shape.crop_top, shape.crop_right, shape.crop_bottom],
                               rotation=shape.rotation)
                    try:
                        batch.add(shape.image.blob, **loc)
                    except Exception:
                        batch.warn("image_extraction_failed", "A slide image could not be recovered.", **loc)
        walk(slide.shapes)
        # Slide background images are not picture shapes.
        for blip in slide._element.findall(".//" + A + "blip"):
            if not any(blip is b for shape in slide.shapes for b in shape._element.iter(A + "blip")):
                rid = blip.get(R + "embed")
                if rid:
                    try:
                        batch.add(slide.part.related_part(rid).blob, slide=number, kind="embedded", role="background")
                    except Exception:
                        batch.warn("image_extraction_failed", "A slide background image could not be recovered.", slide=number)


def docx_images(data, batch):
    with package(data, "docx") as archive:
        parts = ["word/document.xml"] + sorted(n for n in archive.namelist() if
            n.startswith(("word/header", "word/footer")) and n.endswith(".xml"))
        parts += [n for n in ("word/footnotes.xml", "word/endnotes.xml") if n in archive.namelist()]
        for part in parts:
            folder, name = posixpath.split(part)
            relpath = folder + "/_rels/" + name + ".rels"
            if relpath not in archive.namelist():
                continue
            rels = {r.get("Id"): r for r in xml(archive, relpath)}
            root = xml(archive, part)
            caption_paragraphs = {}
            for n, paragraph_node in enumerate(root.iter(W + "p"), 1):
                style = paragraph_node.find("./" + W + "pPr/" + W + "pStyle")
                if style is not None and (style.get(W + "val") or "").casefold() == "caption":
                    caption_paragraphs[n] = "".join(t.text or "" for t in paragraph_node.iter(W + "t"))
            paragraphs = 0
            def walk(node, path, paragraph=None):
                nonlocal paragraphs
                if node.tag == W + "del":
                    return
                if node.tag == W + "p":
                    paragraphs += 1
                    paragraph = paragraphs
                # VML pictures are used by some older Word writers.
                if node.tag in (A + "blip", "{urn:schemas-microsoft-com:vml}imagedata"):
                    if not batch.room():
                        return
                    rid = node.get(R + "embed") or node.get(R + "id")
                    loc = dict(part=part, paragraph=paragraph, xml_path=path)
                    relationship = rels.get(rid)
                    if relationship is None or relationship.get("TargetMode") == "External":
                        batch.warn("external_image_skipped", "External image references are not fetched.", **loc)
                    else:
                        target = relationship.get("Target", "")
                        resolved = posixpath.normpath(posixpath.join(folder, target)) if not target.startswith("/") else target.lstrip("/")
                        if ".." in PurePosixPath(resolved).parts or resolved not in archive.namelist():
                            batch.warn("invalid_image_relationship", "An image relationship could not be resolved.", **loc)
                        else:
                            occurrence = batch.add(archive.read(resolved), **loc)
                            if occurrence and paragraph:
                                occurrence.caption = caption_paragraphs.get(paragraph + 1) or caption_paragraphs.get(paragraph - 1)
                for i, child in enumerate(node, 1):
                    walk(child, path + [i], paragraph)
            walk(root, [])


def extract_images(data, out, settings):
    batch = ImageBatch(out, settings)
    if out.format in {"png", "jpg", "jpeg"}:
        out.page_count, out.page_basis = 1, "standalone_image"
        item = batch.add(data, kind="standalone", page=1)
        if item is None:
            raise ExtractionError("invalid_image", "Image is corrupted, unsupported, or exceeds configured limits.")
        expected = "image/png" if out.format == "png" else "image/jpeg"
        if item.mime_type != expected:
            raise ExtractionError("type_mismatch", "Image contents do not match the filename extension.")
    else:
        extractor = {"pdf": pdf_images, "pptx": pptx_images, "docx": docx_images}.get(out.format)
        if extractor:
            try:
                extractor(data, batch)
            except Exception:
                batch.warn("image_extraction_failed", "Some document images could not be recovered; text remains available.")
    return batch
