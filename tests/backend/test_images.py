"""Image analysis: intake, OCR, visual review, patterns, pipeline and preview."""
from io import BytesIO

import pymupdf
import pytest
from docx import Document as Word
from docx.shared import Inches as WordInches
from openpyxl import Workbook
from openpyxl.drawing.image import Image as SheetImage
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches

from backend.ai.analysis.care_engine import CAREEngine
from backend.ai.llm.analyser import LLMAnalyser
from backend.ai.llm.client import Completion, LLMError
from backend.ai.preprocessing.entity_analyser import NLPAnalyser
from backend.ai.rules.rule_engine import RuleAnalyser
from backend.ai.schemas.concern_schema import ModelAnalysis
from backend.ai.vision.visual_representation import VisionAnalyser, VisualReview, visual_patterns
from backend.extraction.common import ExtractionError
from backend.extraction.image_extractor import FoundImage, collect_images
from backend.extraction.ocr import group_paragraphs, read_text
from backend.services.analysis_service import analyse_extraction
from backend.services.file_service import extract_file

SENTENCE = "Every engineer should check his work."


def picture(text=None, size=(640, 360), colour="white"):
    image = Image.new("RGB", size, colour)
    if text:
        ImageDraw.Draw(image).text((20, 40), text, fill="black", font=ImageFont.truetype("arial.ttf", 28))
    else:
        ImageDraw.Draw(image).rectangle((100, 80, 300, 260), fill="steelblue")
    out = BytesIO()
    image.save(out, "PNG")
    return out.getvalue()


# Intake --------------------------------------------------------------------------
def test_image_upload_is_validated_by_content():
    result = extract_file(picture(), "slide.png")
    assert result.format == "png" and result.page_count == 1 and result.document is None
    with pytest.raises(ExtractionError) as wrong:
        extract_file(picture(), "slide.jpg")
    assert wrong.value.code == "type_mismatch"
    with pytest.raises(ExtractionError) as broken:
        extract_file(b"not an image", "slide.png")
    assert broken.value.code == "invalid_document"


def test_pdf_embedded_and_scanned_pages():
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "A lesson page with a picture.")
    page.insert_image(pymupdf.Rect(72, 100, 372, 280), stream=picture())
    pdf.new_page().insert_image(pymupdf.Rect(0, 0, 595, 842), stream=picture(SENTENCE, (800, 1100)))
    data = pdf.tobytes()
    images, skipped = collect_images(data, extract_file(data, "lesson.pdf").model_dump(mode="json"))
    assert [i.kind for i in images] == ["embedded", "scanned_page"] and not skipped
    embedded = images[0]
    assert embedded.location == {"page": 1} and embedded.label == "Image 1 on page 1"
    assert embedded.bbox[0] == pytest.approx(72 / 595, abs=0.01) and embedded.bbox[2] == pytest.approx(300 / 595, abs=0.01)
    assert images[1].bbox == [0, 0, 1, 1]


def test_office_images_with_locations_skipping_icons_and_repeats():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(BytesIO(picture()), Inches(1), Inches(1), Inches(4), Inches(2.25))
    slide.shapes.add_picture(BytesIO(picture()), Inches(5), Inches(1), Inches(4), Inches(2.25))   # repeat
    slide.shapes.add_picture(BytesIO(picture(size=(20, 20))), Inches(1), Inches(4))               # icon
    buffer = BytesIO()
    deck.save(buffer)
    images, _ = collect_images(buffer.getvalue(), {"format": "pptx"})
    assert len(images) == 1 and images[0].location == {"slide": 1}
    assert images[0].bbox[0] == pytest.approx(1 / 10, abs=0.01)

    doc = Word()
    doc.add_paragraph("Before the picture.")
    doc.add_picture(BytesIO(picture(colour="lightyellow")), width=WordInches(3))
    buffer = BytesIO()
    doc.save(buffer)
    images, _ = collect_images(buffer.getvalue(), {"format": "docx"})
    assert len(images) == 1 and images[0].bbox is None and images[0].label == "Image 1 in the document"

    book = Workbook()
    book.active.add_image(SheetImage(BytesIO(picture(colour="mistyrose"))), "C3")
    buffer = BytesIO()
    book.save(buffer)
    images, _ = collect_images(buffer.getvalue(), {"format": "xlsx"})
    assert len(images) == 1 and images[0].location == {"sheet": "Sheet", "cell": "C3"}


def test_image_limit_is_reported_not_silent():
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 60), "Several pictures.")
    for i, colour in enumerate(["red", "green", "blue"]):
        page.insert_image(pymupdf.Rect(72, 80 + i * 200, 372, 260 + i * 200), stream=picture(colour=colour))
    data = pdf.tobytes()
    images, skipped = collect_images(data, extract_file(data, "lesson.pdf").model_dump(mode="json"), limit=2)
    assert len(images) == 2 and [s["reason"] for s in skipped] == ["image_limit"]


# OCR -----------------------------------------------------------------------------
def test_ocr_reads_text_with_line_boxes():
    paragraphs, low = read_text(picture(SENTENCE))
    assert low == 0 and len(paragraphs) == 1
    assert "engineer should check his work" in paragraphs[0]["text"]
    line = paragraphs[0]["lines"][0]
    assert line["start"] == 0 and all(0 <= v <= 1 for v in line["box"])


def test_wrapped_lines_join_into_one_paragraph_and_gaps_split():
    lines = [{"text": "Every engineer should", "box": [0.1, 0.10, 0.5, 0.05]},
             {"text": "check his work.", "box": [0.1, 0.16, 0.3, 0.05]},
             {"text": "Separate caption", "box": [0.1, 0.60, 0.3, 0.05]},
             {"text": "Second label", "box": [0.6, 0.60, 0.3, 0.05]}]
    paragraphs = group_paragraphs(lines)
    assert [p["text"] for p in paragraphs] == ["Every engineer should check his work.", "Separate caption", "Second label"]
    assert paragraphs[0]["lines"][1] == {"start": 22, "end": 37, "box": [0.1, 0.16, 0.3, 0.05]}


# Visual review -------------------------------------------------------------------
class ScriptedVision:
    def __init__(self, review):
        self.review, self.images = review, []

    def structured(self, prompt, user, schema, sampling=None, images=None):
        assert schema is VisualReview and "never their identity" in prompt and images
        self.images.append(images)
        return VisualReview.model_validate(self.review)


def found(image_id="img1", label="Uploaded image", location=None):
    return FoundImage(image_id, "upload_frame", label, picture(), 640, 360, image_id + "sha", location or {"page": 1}, [0, 0, 1, 1])


REVIEW = {"image_type": "illustration", "summary": "Two people at a desk.", "visible_text": "", "people": [
    {"id": "p1", "presentation": "man", "role": "technical", "activity": "typing code", "prominence": "main"},
    {"id": "p2", "presentation": "woman", "role": "supporting", "activity": "taking notes", "prominence": "secondary"},
    {"id": "p3", "presentation": "unclear", "role": "other", "activity": "walking past", "prominence": "background"}],
    "concerns": [
        {"criterion": "Stereotypical Imagery", "person_ids": ["p1", "p2"], "severity": "Medium", "confidence": "Medium",
         "explanation": "The man codes while the woman takes notes.", "recommendation": "Show both people coding."},
        {"criterion": "Stereotypical Imagery", "person_ids": ["p9"], "severity": "Medium", "confidence": "Low",
         "explanation": "Invented person.", "recommendation": "x"},
        {"criterion": "Technical Confidence Portrayal", "person_ids": ["p3"], "severity": "Low", "confidence": "Low",
         "explanation": "Relies on an unclear person.", "recommendation": "x"},
        {"criterion": "gendered aesthetics / \"pink it and shrink it\"", "person_ids": [], "severity": "Low", "confidence": "Medium",
         "explanation": "Pink is used for the woman's section.", "recommendation": "Use a neutral palette."}]}


def test_vision_keeps_only_grounded_concerns():
    client = ScriptedVision(REVIEW)
    review = VisionAnalyser(client).analyse(found())
    assert [c["criterion"] for c in review["concerns"]] == ["Stereotypical Imagery", "Gendered Aesthetics / “Pink It and Shrink It”"]
    assert review["dropped_concerns"] == 2 and len(client.images[0]) == 1


def test_visual_patterns_need_three_images_and_dominance():
    def review(image_id, presentation):
        return {"image_id": image_id, "status": "analysed", "people": [
            {"id": "p1", "presentation": presentation, "role": "technical", "activity": "coding", "prominence": "main"}]}
    images = [found(f"img{i}", f"Image {i}") for i in range(1, 5)]
    biased = visual_patterns("doc", images, [review(f"img{i}", "man") for i in range(1, 4)])
    assert len(biased) == 1 and biased[0]["scope"] == "document_pattern" and len(biased[0]["images"]) == 3
    balanced = visual_patterns("doc", images, [review("img1", "man"), review("img2", "woman"), review("img3", "man"), review("img4", "woman")])
    assert balanced == []


# Pipeline ------------------------------------------------------------------------
class Quiet:
    provider = "test-only"

    def complete(self, *args, **kwargs):
        return Completion(ModelAnalysis(concerns=[]), "fixture")


def engine():
    return CAREEngine({"rules": RuleAnalyser, "nlp": NLPAnalyser, "llm": lambda: LLMAnalyser(Quiet())})


def test_image_upload_runs_ocr_text_checks_and_visual_review():
    data = picture(SENTENCE)
    extraction = extract_file(data, "worksheet.png")
    stored, progress = [], []
    response = analyse_extraction(extraction, engine=engine(), original=data, on_image=stored.append,
                                  on_progress=lambda *a: progress.append(a),
                                  vision_factory=lambda: VisionAnalyser(ScriptedVision(REVIEW)))
    sections = response["extraction"]["document"]["sections"]
    assert sections[0]["type"] == "ocr_text" and "check his work" in sections[0]["text"]
    assert response["extraction"]["sources"][sections[0]["section_id"]]["kind"] == "image_text"
    findings = response["analysis"]["findings"]
    rule = next(f for f in findings if "Rules" in f["sources"])
    assert rule["evidence"][0]["type"] == "ocr_text" and rule["evidence"][0]["original_content"] == "his"
    visual = [f for f in findings if f["sources"] == ["Vision"]]
    assert {f["criterion"] for f in visual} == {"Stereotypical Imagery", "Gendered Aesthetics / “Pink It and Shrink It”"}
    assert visual[0]["images"][0]["image_id"] == "img1" and visual[0]["observations"]
    report = response["visual_analysis"]
    assert report["status"] == "completed" and report["reviews"][0]["dropped_concerns"] == 2 and report["ocr"]["paragraphs"] == 1
    assert [i.image_id for i in stored] == ["img1"] and progress[0] == (1, 1, 0, 40) and progress[-1][2:] == (40, 95)
    assert response["analysis"]["summary"]["category_counts"]["Visual Representation"] == 2


def test_picture_without_text_and_offline_vision_is_partial_not_clean():
    data = picture()
    def offline():
        raise LLMError("provider_connection", "Cannot reach local Ollama.")
    response = analyse_extraction(extract_file(data, "photo.png"), engine=engine(), original=data, vision_factory=offline)
    assert response["status"] == "failed" and response["analysis"]["findings"] == []
    warning = response["extraction"]["warnings"][-1]
    assert warning["code"] == "image_review_failed" and "Ollama" in warning["message"]


def test_document_warnings_resolved_once_images_are_reviewed():
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "A lesson with a diagram.")
    page.insert_image(pymupdf.Rect(72, 100, 372, 280), stream=picture())
    data = pdf.tobytes()
    extraction = extract_file(data, "lesson.pdf")
    assert any(w["code"] == "visual_content_deferred" for w in extraction.warnings)
    clean = dict(REVIEW, concerns=[])
    response = analyse_extraction(extraction, engine=engine(), original=data, vision_factory=lambda: VisionAnalyser(ScriptedVision(clean)))
    assert not any(w["code"] == "visual_content_deferred" for w in response["extraction"]["warnings"])
    assert response["visual_analysis"]["status"] == "completed"


# Preview and API -----------------------------------------------------------------
def test_preview_highlights_ocr_lines_and_image_concerns(tmp_path):
    from backend.database.database import Store
    from backend.services.preview_service import preview_manifest
    store = Store(tmp_path / "care.sqlite3")
    data = picture(SENTENCE)
    extraction = extract_file(data, "worksheet.png")
    upload = store.add_document(extraction, data, 10)
    job = store.create_job(upload["upload_id"], 8)
    store.claim()
    response = analyse_extraction(extraction, engine=engine(), original=data, on_image=lambda i: store.put_image(job["job_id"], i),
                                  vision_factory=lambda: VisionAnalyser(ScriptedVision(REVIEW)))
    store.update(job["job_id"], result=response)
    store.finish(job["job_id"], "completed", progress=100)
    manifest = preview_manifest(store, job["job_id"])
    assert len(manifest["pages"]) == 1 and not manifest["unmapped_finding_ids"]
    by_kind = {h["section_id"]: h for h in manifest["highlights"]}
    assert by_kind["img1"]["rect"] == [0, 0, 1, 1]
    ocr = next(h for h in manifest["highlights"] if h["section_id"] != "img1")
    assert 0 < ocr["rect"][2] < 1 and ocr["page"] == 1
    assert store.image(job["job_id"], "img1").startswith(b"\x89PNG")


def test_vision_reading_replaces_ocr_that_misses_stylised_words():
    from backend.extraction.ocr import prefer_reading
    ocr = [{"text": "Boys Gdoit BETTER", "lines": [{"start": 0, "end": 17, "box": [0.1, 0.2, 0.8, 0.5]}]}]
    chosen = prefer_reading(ocr, "Boys do it BETTER")
    assert chosen[0]["text"] == "Boys do it BETTER" and len(chosen) == 1
    assert chosen[0]["lines"][0]["end"] == 17 and chosen[0]["lines"][0]["box"] == pytest.approx([0.1, 0.2, 0.8, 0.5])
    assert prefer_reading(ocr, "") == ocr                          # no reading from the vision model
    assert prefer_reading(ocr, "boys gdoit better") == ocr         # OCR already has every word
    long = [{"text": " ".join(["word"] * 50), "lines": []}]
    assert prefer_reading(long, "something else entirely") == long  # long text keeps OCR's exact lines
    assert prefer_reading([], "Boys do it better")[0]["lines"][0]["box"] == [0, 0, 1, 1]
