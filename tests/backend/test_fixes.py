"""Accept-to-fix: sentence units, merged edits, the API flow and corrected exports."""
from io import BytesIO

import pymupdf
import pytest
from docx import Document as Word
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches

from backend.ai.llm.client import LLMError
from backend.config import BackendSettings
from backend.main import create_app
from backend.services.export_service import WRITERS
from backend.services.file_service import extract_file, extract_text
from backend.services.fix_service import sentence_unit
from tests.backend.job_fixtures import fake_worker

TEXT = "Intro line.\nEvery programmer should test his code before the chairman reviews it. Thanks."


class FakeRewriter:
    def rewrite(self, passages, finding):
        return {k: v.replace("his code", "their code").replace("chairman", "chairperson") for k, v in passages.items()}


class OfflineRewriter:
    def rewrite(self, passages, finding):
        raise LLMError("provider_connection", "Cannot reach local Ollama.")


def test_sentence_unit_expands_to_sentence_and_stops_at_line_breaks():
    start = TEXT.index("his")
    left, right = sentence_unit(TEXT, start, start + 3)
    assert TEXT[left:right] == "Every programmer should test his code before the chairman reviews it."
    whole = TEXT.index("Every")
    assert sentence_unit(TEXT, whole, right) == (left, right)  # Evidence already ending in a full stop.
    assert TEXT[slice(*sentence_unit(TEXT, 0, 5))] == "Intro line."


def finding(fid, quote, category="Gendered Language"):
    start = TEXT.index(quote) - TEXT.index("Every")
    return {"finding_id": fid, "category": category, "explanation": "x", "recommendation": "y",
            "evidence": [{"original_content": quote, "location": {"section_id": "s2.1", "start_char": start, "end_char": start + len(quote)}}]}


@pytest.fixture
def app_store(tmp_path):
    app = create_app(BackendSettings(data_dir=tmp_path, cancel_grace_seconds=0.1), worker=fake_worker,
                     rewriter_factory=FakeRewriter)
    with TestClient(app) as client:
        store = app.state.jobs.store
        extraction = extract_text(TEXT)
        upload = store.add_document(extraction, TEXT.encode(), 10)
        # A finished job whose findings point at the second pasted paragraph.
        job = store.create_job(upload["upload_id"], 8)
        store.claim()
        sid = extraction.document.sections[1].section_id
        result = {"status": "completed", "analysis": {"findings": [finding("f-his", "his"), finding("f-chair", "chairman"),
                                                                    finding("f-stereo", "Every programmer", "Gender Stereotypes")]},
                  "analysis_sources": {"s2.1": {"original_section_id": sid, "source_start_char": 0}}}
        store.update(job["job_id"], result=result)
        store.finish(job["job_id"], "completed", progress=100)
        yield client, store, job["job_id"]


def test_propose_apply_merge_undo_and_download(app_store):
    client, store, job_id = app_store
    proposal = client.post(f"/api/fixes/{job_id}/f-his/proposal").json()
    assert proposal["ai_rewrite"] and len(proposal["units"]) == 1
    unit = proposal["units"][0]
    assert unit["original"] == "Every programmer should test his code before the chairman reviews it."
    edits = [{k: unit[k] for k in ("section_id", "start", "end")} | {"base": unit["current"], "replacement": unit["proposed"]}]
    applied = client.put(f"/api/fixes/{job_id}/f-his", json={"edits": edits})
    assert applied.status_code == 200, applied.text
    assert applied.json()["edits"][0]["finding_ids"] == ["f-his"]
    # A fixed finding cannot be dismissed until its fix is undone.
    assert client.put(f"/api/feedback/{job_id}/f-his", json={"decision": "dismissed"}).status_code == 409

    # A second finding on the same sentence starts from the already-corrected wording.
    second = client.post(f"/api/fixes/{job_id}/f-chair/proposal").json()["units"][0]
    assert "their code" in second["current"] and "chairperson" in second["proposed"]
    merged = client.put(f"/api/fixes/{job_id}/f-chair", json={"edits": [{k: second[k] for k in ("section_id", "start", "end")} | {"base": second["current"], "replacement": second["proposed"]}]}).json()
    assert len(merged["edits"]) == 1 and merged["edits"][0]["finding_ids"] == ["f-chair", "f-his"]
    assert {f["finding_id"] for f in merged["feedback"]} == {"f-chair", "f-his"}

    stale = client.put(f"/api/fixes/{job_id}/f-his", json={"edits": edits})
    assert stale.status_code == 409 and stale.json()["error"]["code"] == "fix_outdated"

    download = client.get(f"/api/results/{job_id}/document")
    assert download.status_code == 200 and download.headers["x-care-unapplied"] == "0"
    assert download.content.decode() == TEXT.replace("his code", "their code").replace("chairman", "chairperson")
    assert "revised" in download.headers["content-disposition"]
    assert client.get(f"/api/results/{job_id}").json()["edits"]

    undone = client.delete(f"/api/fixes/{job_id}/f-his").json()
    assert undone["edits"] == [] and undone["feedback"] == []
    assert client.get(f"/api/results/{job_id}/document").content.decode() == TEXT


def test_unchanged_wording_and_offline_model(app_store, monkeypatch):
    client, store, job_id = app_store
    client.app.state.rewriter_factory = OfflineRewriter
    proposal = client.post(f"/api/fixes/{job_id}/f-stereo/proposal").json()
    assert not proposal["ai_rewrite"] and "Ollama" in proposal["message"]
    unit = proposal["units"][0]
    assert unit["proposed"] == unit["current"]  # The educator types the fix instead.
    same = client.put(f"/api/fixes/{job_id}/f-stereo", json={"edits": [{k: unit[k] for k in ("section_id", "start", "end")} | {"base": unit["current"], "replacement": unit["current"]}]})
    assert same.status_code == 422 and same.json()["error"]["code"] == "fix_unchanged"
    assert client.post(f"/api/fixes/{job_id}/unknown/proposal").status_code == 404


def export(data, name, edits_by_text):
    """Extract, edit the sections containing each phrase, write, and re-extract."""
    extraction = extract_file(data, name)
    sections = {s.section_id: s.text for s in extraction.document.sections}
    by_section = {}
    for old, new in edits_by_text.items():
        sid = next(k for k, v in sections.items() if old in v)
        start = sections[sid].index(old)
        by_section.setdefault(sid, []).append({"section_id": sid, "start": start, "end": start + len(old), "original": old, "replacement": new, "finding_ids": ["f"]})
    output, unapplied = WRITERS[extraction.format](data, extraction.model_dump(mode="json"), sections, by_section)
    assert unapplied == 0
    return [s.text for s in extract_file(output, name).document.sections]


def test_word_fix_keeps_formatting_runs_and_other_text():
    doc = Word()
    p = doc.add_paragraph("Every programmer should test ")
    p.add_run("his").bold = True
    p.add_run(" code.")
    doc.add_paragraph("Untouched paragraph.")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "A good chairman listens."
    buffer = BytesIO()
    doc.save(buffer)
    texts = export(buffer.getvalue(), "lesson.docx", {"Every programmer should test his code.": "Every programmer should test their code.",
                                                     "chairman": "chairperson"})
    assert "Every programmer should test their code." in texts
    assert "Untouched paragraph." in texts and "A good chairperson listens." in texts


def test_powerpoint_fix_in_text_box_and_notes():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = "Hey guys, every engineer checks his work."
    slide.notes_slide.notes_text_frame.text = "Ask the chairman."
    buffer = BytesIO()
    deck.save(buffer)
    texts = export(buffer.getvalue(), "slides.pptx", {"his work": "their work", "chairman": "chairperson"})
    assert "Hey guys, every engineer checks their work." in texts and "Ask the chairperson." in texts


def test_excel_fix_does_not_change_other_cells_sharing_the_string():
    book = Workbook()
    book.active["A1"] = "The chairman decides."
    book.active["B1"] = "The chairman decides."
    buffer = BytesIO()
    book.save(buffer)
    extraction = extract_file(buffer.getvalue(), "sheet.xlsx")
    sid = next(s.section_id for s in extraction.document.sections if extraction.sources[s.section_id].get("cell") == "A1")
    edit = {"section_id": sid, "start": 4, "end": 12, "original": "chairman", "replacement": "chairperson", "finding_ids": ["f"]}
    output, unapplied = WRITERS["xlsx"](buffer.getvalue(), extraction.model_dump(mode="json"),
                                        {s.section_id: s.text for s in extraction.document.sections}, {sid: [edit]})
    sheet = load_workbook(BytesIO(output)).active
    assert unapplied == 0 and sheet["A1"].value == "The chairperson decides." and sheet["B1"].value == "The chairman decides."


def test_pdf_fix_replaces_text_in_place():
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_textbox(pymupdf.Rect(72, 72, 520, 140), "Every programmer should test his code.", fontsize=12)
    page.insert_textbox(pymupdf.Rect(72, 200, 520, 260), "Second block stays.", fontsize=12)
    data = pdf.tobytes()
    texts = " ".join(export(data, "notes.pdf", {"his code": "their code"}))
    assert "their code" in texts and "his code" not in texts and "Second block stays." in texts


def test_plain_text_and_csv_fixes():
    assert export(b"Ask the chairman.\n\nSecond paragraph.", "notes.txt", {"chairman": "chairperson"})[0].startswith("Ask the chairperson.")
    texts = export(b'Name,Note\nA,"Every man for himself"\n', "table.csv", {"Every man for himself": "Everyone for themselves"})
    assert "Everyone for themselves" in texts and "Name" in texts


def test_display_spans_move_highlights_after_fixes():
    from backend.services.preview_service import display_spans
    text = "Ask the chairman. Every programmer checks his code. The men lead."
    sections = {"s1": text}
    result = {"analysis_sources": {"s1.1": {"original_section_id": "s1", "source_start_char": 0}}}
    def ev(fid, quote):
        start = text.index(quote)
        return {"finding_id": fid, "evidence": [{"original_content": quote, "location": {"section_id": "s1.1", "start_char": start, "end_char": start + len(quote)}}]}
    findings = [ev("chair", "chairman"), ev("his", "his"), ev("men", "men lead"), ev("other", "the chairman")]
    edit = {"section_id": "s1", "start": text.index("Ask"), "end": text.index("man.") + 4, "original": "Ask the chairman.",
            "replacement": "Ask the chairperson.", "finding_ids": ["chair"]}
    texts, spans, unmapped = display_spans(result, sections, findings, [edit])
    shown = texts["s1"]
    assert shown.startswith("Ask the chairperson.")
    places = {fid: shown[a:b] for fid, _, a, b in spans}
    # "other" pointed at words the chairman fix rewrote: it is shown on the corrected sentence.
    assert places == {"chair": "Ask the chairperson.", "his": "his", "men": "men lead", "other": "Ask the chairperson."} and not unmapped


def test_pdf_preview_uses_corrected_copy(app_store, tmp_path):
    from backend.services.preview_service import job_page_image, preview_manifest
    client, store, _ = app_store
    pdf = pymupdf.open()
    pdf.new_page().insert_textbox(pymupdf.Rect(72, 72, 520, 140), "Every programmer should test his code.", fontsize=12)
    data = pdf.tobytes()
    extraction = extract_file(data, "notes.pdf")
    upload = store.add_document(extraction, data, 10)
    job = store.create_job(upload["upload_id"], 8)
    store.claim()
    sid = extraction.document.sections[0].section_id
    text = extraction.document.sections[0].text
    start = text.index("his")
    result = {"status": "completed", "analysis": {"findings": [{"finding_id": "f1", "category": "Gendered Language", "evidence": [
        {"original_content": "his", "location": {"section_id": f"{sid}.1", "start_char": start, "end_char": start + 3}}]}]},
        "analysis_sources": {f"{sid}.1": {"original_section_id": sid, "source_start_char": 0}}}
    store.update(job["job_id"], result=result)
    store.finish(job["job_id"], "completed", progress=100)
    before = preview_manifest(store, job["job_id"])
    assert before["revision"] == "original" and before["highlights"]
    sentence = text.strip()
    store.apply_fix(job["job_id"], "f1", [], [{"section_id": sid, "start": text.index(sentence), "end": text.index(sentence) + len(sentence),
                                               "original": sentence, "replacement": "Every programmer should test their code.", "finding_ids": ["f1"]}],
                    {"decision": "accepted", "fix_applied": True})
    after = preview_manifest(store, job["job_id"])
    assert after["revision"] != "original" and [h["finding_id"] for h in after["highlights"]] == ["f1"]
    assert job_page_image(store, job["job_id"], 1).startswith(b"\x89PNG")


class UnchangedRewriter:
    def rewrite(self, passages, finding):
        raise LLMError("rewrite_unchanged", "The local AI did not produce different wording for this concern")


def test_concern_already_corrected_by_another_fix_can_be_marked_fixed(app_store):
    client, store, job_id = app_store
    unit = client.post(f"/api/fixes/{job_id}/f-his/proposal").json()["units"][0]
    client.put(f"/api/fixes/{job_id}/f-his", json={"edits": [{k: unit[k] for k in ("section_id", "start", "end")} | {"base": unit["current"], "replacement": unit["proposed"]}]})
    client.app.state.rewriter_factory = UnchangedRewriter
    proposal = client.post(f"/api/fixes/{job_id}/f-stereo/proposal").json()
    later = proposal["units"][0]
    assert later["already_fixed"] and "already includes a fix" in proposal["message"]
    applied = client.put(f"/api/fixes/{job_id}/f-stereo", json={"edits": [{k: later[k] for k in ("section_id", "start", "end")} | {"base": later["current"], "replacement": later["current"]}]})
    assert applied.status_code == 200 and applied.json()["edits"][0]["finding_ids"] == ["f-his", "f-stereo"]
