from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
import pymupdf
from docx import Document as Word
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from backend.extraction.common import ExtractionError, MAX_BYTES
from backend.services.file_service import extract_file, extract_text


def saved(document):
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_docx_body_table_header_footer_and_offsets():
    word = Word()
    word.add_paragraph("Before the table.")
    word.add_table(1, 2).cell(0, 0).text = "Every engineer should check his work."
    word.tables[0].cell(0, 1).text = "Second cell"
    word.add_paragraph("After the table.")
    word.sections[0].header.paragraphs[0].text = "Teaching header"
    word.sections[0].footer.paragraphs[0].text = "Teaching footer"
    result = extract_file(saved(word), "../lesson.docx")
    texts = [s.text for s in result.document.sections]
    assert texts[:4] == ["Before the table.", "Every engineer should check his work.", "Second cell", "After the table."]
    assert {"Teaching header", "Teaching footer"} <= set(texts)
    assert result.filename == "lesson.docx"
    assert result.sources["s2"]["table"] == 1
    assert result.sources["s2"]["row"] == 1
    assert result.sources["s2"]["cell"] == 1
    assert result.document.sections[0].page is None
    assert result.page_basis.startswith("generated")


def test_pdf_pages_blocks_password_and_scans():
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Every programmer should test his code.")
    pdf.new_page()
    result = extract_file(pdf.tobytes(), "lesson.pdf")
    assert result.page_count == 2
    assert result.document.sections[0].page == 1
    assert len(result.sources["s1"]["bbox"]) == 4
    assert result.status == "partial"
    assert any(w["code"] == "page_without_text" and w["page"] == 2 for w in result.warnings)
    encrypted = pdf.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="secret", owner_pw="owner")
    with pytest.raises(ExtractionError, match="Password"):
        extract_file(encrypted, "locked.pdf")
    blank = pymupdf.open()
    blank.new_page()
    assert extract_file(blank.tobytes(), "scan.pdf").status == "no_text"
    pdf.close()
    blank.close()


def test_pptx_shapes_tables_notes_and_reading_order():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1)).text = "Second paragraph"
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = "First paragraph"
    slide.shapes.add_table(1, 1, Inches(1), Inches(3), Inches(3), Inches(1)).table.cell(0, 0).text = "Cell content"
    slide.notes_slide.notes_text_frame.text = "Speaker note"
    result = extract_file(saved(deck), "lesson.pptx")
    assert [s.text for s in result.document.sections][:4] == ["First paragraph", "Second paragraph", "Cell content", "Speaker note"]
    assert all(s.slide == 1 for s in result.document.sections)
    assert result.page_basis == "native_slides"


@pytest.mark.parametrize("extension", ["xlsx", "xlsm"])
def test_excel_text_coordinates_hidden_sheets_and_uncached_formulas(extension):
    book = Workbook()
    sheet = book.active
    sheet.title = "Teaching"
    sheet["C8"] = "Women are naturally better at communication."
    sheet["A1"] = "=1+2"
    hidden = book.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["A2"] = "Hidden teaching text"
    result = extract_file(saved(book), "lesson." + extension)
    content = next(s for s in result.document.sections if "Women" in s.text)
    assert result.sources[content.section_id]["cell"] == "C8"
    assert result.sources[content.section_id]["sheet"] == "Teaching"
    assert any(s.text == "Hidden teaching text" for s in result.document.sections)
    assert any(w["code"] == "uncached_formula" for w in result.warnings)
    assert not any(s.text == "=1+2" for s in result.document.sections)


@pytest.mark.parametrize("name,data,expected", [
    ("notes.txt", "English — text".encode(), "English — text"),
    ("notes.md", b"# Teaching\r\nA paragraph.", "# Teaching\r\nA paragraph."),
    ("notes.txt", "UTF16 text".encode("utf-16"), "UTF16 text"),
    ("notes.csv", b'Name,Description\nExample,"line one\nline two"', "line one\nline two"),
    ("notes.tsv", b"Name\tDescription\nExample\tTeaching text", "Teaching text"),
])
def test_plain_formats(name, data, expected):
    result = extract_file(data, name)
    assert expected in [s.text for s in result.document.sections]


@pytest.mark.parametrize("data,name,code", [
    (b"", "a.txt", "empty_file"), (b"hello", "a.exe", "unsupported_format"),
    (b"not PDF", "a.pdf", "type_mismatch"), (b"broken zip", "a.docx", "invalid_document"),
    (b"\x00binary", "a.txt", "binary_text"), (b"\xffbad", "a.txt", "text_encoding"),
    (bytes.fromhex("D0CF11E0A1B11AE1"), "a.docx", "encrypted_or_legacy_office"),
])
def test_invalid_inputs(data, name, code):
    with pytest.raises(ExtractionError) as error:
        extract_file(data, name)
    assert error.value.code == code


def test_size_and_page_limits():
    with pytest.raises(ExtractionError, match="20 MB"):
        extract_file(b"a" * (MAX_BYTES + 1), "big.txt")
    with pytest.raises(ExtractionError, match="100"):
        extract_text("a" * 300001)
    pdf = pymupdf.open()
    for _ in range(101):
        pdf.new_page()
    with pytest.raises(ExtractionError, match="100"):
        extract_file(pdf.tobytes(), "big.pdf")
    pdf.close()


def test_mismatched_office_and_unsafe_xml():
    with pytest.raises(ExtractionError) as error:
        extract_file(saved(Word()), "renamed.xlsx")
    assert error.value.code == "type_mismatch"
    stream = BytesIO()
    with ZipFile(stream, "w") as archive:
        archive.writestr("[Content_Types].xml", "<types/>")
        archive.writestr("word/document.xml", '<!DOCTYPE x [<!ENTITY data "EXPANSION">]><x>&data;</x>')
    with pytest.raises(ExtractionError):
        extract_file(stream.getvalue(), "unsafe.docx")


def test_real_legacy_xls_fixture():
    path = Path(__file__).resolve().parents[2] / "sample_data/text/legacy_example.xls"
    result = extract_file(path.read_bytes(), path.name)
    section = next(s for s in result.document.sections if "programmer" in s.text)
    assert result.sources[section.section_id]["cell"] == "C2"
    assert result.sources[section.section_id]["sheet"] == "Teaching"


def test_zip_traversal_is_rejected_without_unpacking():
    stream = BytesIO()
    with ZipFile(stream, "w") as archive:
        archive.writestr("../outside.xml", "<x/>")
    with pytest.raises(ExtractionError) as error:
        extract_file(stream.getvalue(), "unsafe.docx")
    assert error.value.code == "invalid_archive"


def test_docx_hyperlink_and_textbox_text_is_not_lost_or_duplicated():
    from docx.oxml import parse_xml
    word = Word()
    paragraph = word.add_paragraph("Main paragraph. ")
    paragraph._p.append(parse_xml('<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:r><w:t>Linked text</w:t></w:r></w:hyperlink>'))
    paragraph._p.append(parse_xml('<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:pict><w:txbxContent><w:p><w:r><w:t>Text box content</w:t></w:r></w:p></w:txbxContent></w:pict></w:r>'))
    result = extract_file(saved(word), "boxes.docx")
    assert [s.text for s in result.document.sections] == ["Main paragraph. Linked text", "Text box content"]


def test_pdf_table_layout_retained_without_duplicate_text():
    pdf = pymupdf.open()
    page = pdf.new_page()
    for x in (50, 200, 350):
        page.draw_line((x, 50), (x, 150))
    for y in (50, 100, 150):
        page.draw_line((50, y), (350, y))
    for x, y, text in [(60, 75, "Role"), (210, 75, "Example"), (60, 125, "Engineer"), (210, 125, "Teaching")]:
        page.insert_text((x, y), text)
    result = extract_file(pdf.tobytes(), "table.pdf")
    tables = [t for source in result.sources.values() for t in source["tables"]]
    assert tables and tables[0]["rows"] == 2 and tables[0]["columns"] == 2
    assert "".join(s.text for s in result.document.sections).count("Engineer") == 1
    pdf.close()


def test_pasted_paragraph_locations_and_unicode_are_exact():
    text = "First paragraph.\r\n\r\nSecond paragraph — with emoji 🙂."
    result = extract_text(text)
    assert len(result.document.sections) == 2
    for section in result.document.sections:
        source = result.sources[section.section_id]
        assert text[source["input_start_char"]:source["input_end_char"]] == section.text


def test_pasted_text_line_breaks_are_paragraphs():
    text = "First paragraph.\nSecond paragraph.\r\n\r\nThird paragraph."
    result = extract_text(text)
    sections = result.document.sections
    assert [s.text.strip() for s in sections] == ["First paragraph.", "Second paragraph.", "Third paragraph."]
    assert [s.paragraph for s in sections] == [1, 2, 3]
    assert "".join(s.text for s in sections) == text
    for s in sections:
        source = result.sources[s.section_id]
        assert text[source["input_start_char"]:source["input_end_char"]] == s.text
