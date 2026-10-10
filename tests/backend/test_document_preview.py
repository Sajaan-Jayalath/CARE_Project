from io import BytesIO
from concurrent.futures import ThreadPoolExecutor
import pytest
import pymupdf
from docx import Document

from backend.database.database import Store, BackendError
from backend.services.file_service import extract_file
from backend.services.preview_service import preview_manifest, preview_pdf, page_image, office_executable
from fastapi.testclient import TestClient
from types import SimpleNamespace
from backend.main import create_app
from backend.config import BackendSettings


def setup_result(tmp_path, data, name, quote="his"):
    store = Store(tmp_path / "preview.sqlite3")
    extraction = extract_file(data, name)
    upload = store.add_document(extraction, data, 10)
    job = store.create_job(upload["upload_id"], 10)
    section = next(s for s in extraction.document.sections if quote in s.text)
    start = section.text.index(quote)
    finding = {"finding_id": "f1", "severity": "High", "category": "Gendered Language", "evidence": [{"original_content": quote,
        "location": {"section_id": "segment", "start_char": start, "end_char": start + len(quote)}}]}
    result = {"analysis_sources": {"segment": {"original_section_id": section.section_id, "source_start_char": 0}},
              "analysis": {"findings": [finding]}}
    store.update(job["job_id"], status="completed", result=result)
    return store, upload, job, result


def pdf_bytes(rotation=0, repeat=False):
    with pymupdf.open() as pdf:
        p = pdf.new_page()
        p.insert_text((72, 80), "Every programmer should test his code.")
        if repeat:
            p.insert_text((72, 150), "Every programmer should test his code.")
        p.set_rotation(rotation)
        return pdf.tobytes()


@pytest.mark.parametrize("rotation", [0, 90, 180, 270])
def test_original_pdf_highlight_and_rotation(tmp_path, rotation):
    store, upload, job, _ = setup_result(tmp_path, pdf_bytes(rotation, repeat=True), "lesson.pdf")
    manifest = preview_manifest(store, job["job_id"])
    assert not manifest["converted"] and not manifest["unmapped_finding_ids"]
    assert len(manifest["highlights"]) == 1  # repeated paragraph distinguished by source bbox
    rect = manifest["highlights"][0]["rect"]
    assert all(0 <= v <= 1 for v in rect)
    assert page_image(store, upload["upload_id"], 1).startswith(b"\x89PNG")
    with pytest.raises(BackendError, match="does not exist"):
        page_image(store, upload["upload_id"], 2)


def test_mismatched_evidence_is_not_highlighted(tmp_path):
    store, _, job, result = setup_result(tmp_path, pdf_bytes(), "lesson.pdf")
    result["analysis"]["findings"][0]["evidence"][0]["original_content"] = "her"
    store.update(job["job_id"], result=result)
    manifest = preview_manifest(store, job["job_id"])
    assert manifest["highlights"] == []
    assert manifest["unmapped_finding_ids"] == ["f1"]


def test_preview_api_returns_manifest_and_original_page_image(tmp_path):
    store, upload, job, _ = setup_result(tmp_path, pdf_bytes(), "lesson.pdf")
    app = create_app(BackendSettings(data_dir=tmp_path))
    app.state.jobs = SimpleNamespace(store=store)
    client = TestClient(app)
    response = client.get(f'/api/results/{job["job_id"]}/preview')
    assert response.status_code == 200 and response.json()["highlights"]
    image = client.get(f'/api/uploads/{upload["upload_id"]}/preview/pages/1')
    assert image.status_code == 200 and image.headers["content-type"] == "image/png"
    assert image.headers["cache-control"] == "no-store"
    assert client.get('/api/results/missing/preview').status_code == 404


@pytest.mark.parametrize("fmt", ["pptx", "xlsx"])
@pytest.mark.skipif(not office_executable(), reason="Requires local LibreOffice")
def test_slides_and_spreadsheets_render_locally(tmp_path, fmt):
    stream = BytesIO()
    if fmt == "pptx":
        from pptx import Presentation
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = "Lesson"
        slide.placeholders[1].text = "Every programmer should test his code."
        deck.save(stream)
    else:
        from openpyxl import Workbook
        book = Workbook()
        book.active['A1'] = "Every programmer should test his code."
        book.active.column_dimensions['A'].width = 65
        book.save(stream)
    store, _, job, _ = setup_result(tmp_path, stream.getvalue(), f"lesson.{fmt}")
    manifest = preview_manifest(store, job["job_id"])
    assert manifest["converted"] and manifest["highlights"]
    assert not manifest["unmapped_finding_ids"]


@pytest.mark.skipif(not office_executable(), reason="Requires local LibreOffice")
def test_office_original_layout_conversion_cache_and_delete(tmp_path):
    doc = Document()
    doc.add_heading("Teaching Material", 0)
    doc.add_paragraph("Every programmer should test his code.")
    doc.add_table(rows=2, cols=2).cell(0, 0).text = "Table preserved"
    stream = BytesIO(); doc.save(stream)
    store, upload, job, _ = setup_result(tmp_path, stream.getvalue(), "lesson.docx")
    with ThreadPoolExecutor(max_workers=2) as pool:
        previews = list(pool.map(lambda _: preview_manifest(store, job["job_id"]), range(2)))
    manifest = previews[0]
    assert previews[0] == previews[1]
    assert manifest["converted"] and len(manifest["pages"]) == 1
    assert len(manifest["highlights"]) == 1
    assert not manifest["unmapped_finding_ids"]
    document = store.document(upload["upload_id"], original=True)
    assert preview_pdf(store, document).startswith(b"%PDF")
    with store.connection() as db:
        assert db.execute("SELECT COUNT(*) FROM document_previews").fetchone()[0] == 1
    store.delete_document(upload["upload_id"])
    with store.connection() as db:
        assert db.execute("SELECT COUNT(*) FROM document_previews").fetchone()[0] == 0
